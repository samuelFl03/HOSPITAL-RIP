from docx import Document
from docx.shared import Inches, Pt, RGBColor
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(0.8)
for name in ['Normal','Title','Heading 1']:
 doc.styles[name].font.name='Calibri'
 doc.styles[name].font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.add_heading('Hospital RIP Nómina por horas',0)
doc.add_paragraph('Actualización versión 1.4 | 9 de octubre de 2026')
doc.add_paragraph('Esta versión incorpora trabajadores, jornadas trabajadas y liquidaciones de salario por horas. El administrador gestiona la nómina desde Nómina y pagos y conserva el historial de cada pago.')
doc.add_heading('Cómo registrar y liquidar',1)
for line in ['Registrar trabajador: ingresar nombre, identificación, cargo y tarifa por hora en pesos colombianos.','Registrar jornada: seleccionar trabajador y registrar inicio y fin en hora de Colombia. Solo se admiten jornadas ya realizadas, sin cruces y de hasta 24 horas.','Liquidar horas: seleccionar trabajador y período inclusive. Se incluyen jornadas completas pendientes; el salario es la suma de horas por la tarifa de cada jornada, redondeada al centavo.','Registrar pago: después de realizar el pago, ingresar la referencia del comprobante. La aplicación registra el pago; no realiza transferencias bancarias.','Consultar historial: revisar liquidaciones pendientes o pagadas, exportar CSV o imprimir desde el navegador.']:
 doc.add_paragraph(line,style='List Number')
doc.add_heading('Conservación y permisos',1)
doc.add_paragraph('Cada jornada conserva la tarifa registrada aunque se cambie la tarifa del trabajador. Cada liquidación conserva nombre e identificación. Las jornadas liquidadas no se retiran y las horas no pueden liquidarse dos veces. El servidor valida permisos, fechas y tarifas y registra las operaciones en auditoría. La nómina es exclusiva del administrador.')
doc.add_heading('Alcance del salario',1)
doc.add_paragraph('Una jornada de 8 horas con tarifa de 25 000 COP produce un salario de 200 000 COP. No se agregan salario fijo, recargos ni descuentos automáticamente. Las guardias planificadas no generan horas trabajadas de forma automática.')
doc.add_heading('Instalación y comprobaciones',1)
doc.add_paragraph('El esquema incluye workers, work_hours y payrolls. Para una instancia local existente ejecutar node scripts/migrate-payroll.js con MySQL iniciado. Para instalaciones nuevas usar el procedimiento de configuración habitual. Se mantienen las tablas y los registros anteriores.')
doc.add_paragraph('Pruebas: npm test valida tarifas, redondeo y permisos. npm run test:integration verifica el flujo completo en una base temporal y conserva hospital_rip. Los resultados de la ejecución se informan junto con la entrega.')
doc.save('docs/Hospital_RIP_Actualizacion_v1_4.docx')
