from docx import Document
from docx.shared import Inches, Pt, RGBColor
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(0.8)
for name in ['Normal','Title','Heading 1']:
 doc.styles[name].font.name='Calibri'
 doc.styles[name].font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.add_heading('Hospital RIP Selección de cargos',0)
doc.add_paragraph('Actualización versión 1.7 | 9 de octubre de 2026')
doc.add_paragraph('El registro y la edición de trabajadores incorporan una lista de cargos disponibles. La selección es obligatoria y el servidor valida que el cargo pertenezca a las opciones enviadas por la API. Los cargos existentes se conservan.')
doc.add_heading('Cargos disponibles',1)
doc.add_paragraph('Administrador, Auxiliar, Auxiliar de enfermería, Camillero, Enfermero, Médico general, Médico especialista, Médico jefe, Personal de laboratorio, Personal de limpieza, Personal de mantenimiento, Psicólogo, Recepcionista y Supervisor. También aparecen los cargos ya registrados en trabajadores, sin duplicados y ordenados alfabéticamente.')
doc.add_heading('Uso y permisos',1)
doc.add_paragraph('En Nómina y pagos, abrir Registrar trabajador, desplegar Cargo y seleccionar una opción. Al editar un trabajador se muestra su cargo actual. El cargo laboral no cambia los permisos de acceso: estos corresponden al rol de la cuenta de usuario.')
doc.add_heading('Actualización y validación',1)
doc.add_paragraph('Subir los archivos de esta versión a GitHub y esperar el despliegue en Railway. Recargar la aplicación. No requiere modificar tablas ni importar nuevamente los datos. Las pruebas comprueban la conservación de cargos existentes, la lista de opciones y el rechazo de cargos inválidos; la integración se ejecuta en una base temporal.')
doc.save('docs/Hospital_RIP_Actualizacion_v1_7.docx')
