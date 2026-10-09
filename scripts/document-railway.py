from docx import Document
from docx.shared import Inches, Pt, RGBColor
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(0.8)
for name in ['Normal','Title','Heading 1']:
 doc.styles[name].font.name='Calibri'
 doc.styles[name].font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.add_heading('Hospital RIP Instalación en Railway',0)
doc.add_paragraph('Actualización versión 1.5 | 9 de octubre de 2026')
doc.add_paragraph('Esta actualización corrige la instalación de tablas en Railway. El instalador utiliza la base indicada por DB_NAME y conserva los datos existentes. El error ER_NO_SUCH_TABLE durante el inicio de sesión indica que faltan tablas en la base utilizada por la aplicación.')
doc.add_heading('Preparación',1)
doc.add_paragraph('Configurar DB_HOST, DB_PORT, DB_NAME, DB_USER y DB_PASSWORD con referencias al servicio MySQL de Railway. Mantener HOST como 0.0.0.0, NODE_ENV como production y APP_ORIGIN como el dominio HTTPS público sin barra final.')
doc.add_paragraph('Para una base sin usuarios, definir INITIAL_ADMIN_EMAIL y INITIAL_ADMIN_PASSWORD. La contraseña debe contener entre 10 y 128 caracteres. INITIAL_ADMIN_NAME es opcional. Estos datos se guardan en las variables privadas de Railway, fuera del repositorio.')
doc.add_heading('Instalación y acceso',1)
doc.add_paragraph('Subir la actualización al repositorio conectado. En Settings y Deploy del servicio HOSPITAL-RIP, establecer Pre-Deploy Command como npm run setup. Mantener Start Command como npm start. Aplicar cambios y revisar que la instalación termine correctamente antes de iniciar sesión.')
doc.add_paragraph('El instalador crea las tablas de gestión hospitalaria y nómina. Solo crea el administrador inicial cuando no hay usuarios; conserva las cuentas existentes y no carga datos ficticios en producción. Iniciar sesión con los datos configurados. Después pueden retirarse las variables de la cuenta inicial.')
doc.add_heading('Validación',1)
doc.add_paragraph('Las pruebas verifican la selección de una base distinta de hospital_rip y el rechazo de nombres inválidos. La integración utiliza una base temporal para comprobar los flujos y conserva hospital_rip. Los registros de Railway permiten confirmar que la instalación se ejecutó en el despliegue real.')
doc.save('docs/Hospital_RIP_Actualizacion_v1_5.docx')
